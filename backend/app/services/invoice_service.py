from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.invoice import (
    ChargePlan, ChargePlanRate, CountryResolutionRule, FinalInvoiceLine,
    InvoiceException, InvoiceProcess, LocationCountryMapping, Organisation,
    ScmUsage, ShipmentData, UNLOCODE
)
from app.services.edc_client import read_shipments, read_usage

GENERATED_DIR = Path("/app/generated")


def parse_dt(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except Exception:
        return None


def upsert_organisation(db: Session, org_code: str, org_full_name: Optional[str]):
    org = db.query(Organisation).filter(Organisation.org_code == org_code).first()
    if org is None:
        org = Organisation(org_code=org_code, org_full_name=org_full_name, charge_plan_code="DEFAULT_2026")
        db.add(org)
    else:
        org.org_full_name = org_full_name or org.org_full_name
        org.updated_at = datetime.utcnow()
    if org.get_shipment_data:
        org.country_code = "Multi"
    return org


def make_usage(process_number: int, row: Dict[str, Any]) -> ScmUsage:
    return ScmUsage(
        process_number=process_number,
        org_code=row.get("OrgCode"),
        org_full_name=row.get("OrgFullName"),
        year_month=row.get("YearMonth"),
        shipment_count=row.get("ShipmentCount") or 0,
        order_count=row.get("OrderCount") or 0,
        pos=row.get("POs") or 0,
        stos=row.get("STOs") or 0,
        sos=row.get("SOs") or 0,
        clos=row.get("CLOs") or 0,
        order_edi_count=row.get("OrderEDICount") or 0,
        order_from_shipment_edi_count=row.get("OrderFromShipmentEDICount") or 0,
        order_api_count=row.get("OrderAPICount") or 0,
        shipment_edi_count=row.get("ShipmentEDICount") or 0,
        invoice_count=row.get("InvoiceCount") or 0,
        booking_count=row.get("BookingCount") or 0,
        orders_shipped=row.get("OrdersShipped") or 0,
        shipment_event_edi_count=row.get("ShipmentEventEDICount") or 0,
        first_order_date=parse_dt(row.get("FirstOrderDate")),
        last_order_date=parse_dt(row.get("LastOrderDate")),
        first_shipment_date=parse_dt(row.get("FirstShipmentDate")),
        last_shipment_date=parse_dt(row.get("LastShipmentDate")),
        org_date_created=parse_dt(row.get("OrgDateCreated")),
        org_managed_by=row.get("OrgManagedBy"),
        org_active=row.get("OrgActive"),
        is_road_provider=row.get("IsRoadProvider"),
        is_air_sea_provider=row.get("IsAirSeaProvider"),
        is_solutions_provider=row.get("IsSolutionsProvider"),
        is_sales_order_client=row.get("IsSalesOrderClient"),
        is_stock_transfer_order_client=row.get("IsStockTransferOrderClient"),
        is_llp=row.get("IsLLP"),
        is_edc_lite_client=row.get("IsEdcLiteClient"),
        is_freight_manager=row.get("IsFreightManager"),
        is_multi_lsppom=row.get("IsMultiLSPPOM"),
        is_po_carrier_allocation_manager=row.get("IsPOCarrierAllocationManager"),
        is_po_control_client=row.get("IsPOControlClient"),
        is_purchase_order_client=row.get("IsPurchaseOrderClient"),
        is_po_visibility_client=row.get("IsPOVisibilityClient"),
        is_po_manager_client=row.get("IsPOManagerClient"),
        is_supply_chain_manager=row.get("IsSupplyChainManager"),
        is_supply_chain_optimizer=row.get("IsSupplyChainOptimizer"),
        is_freight_order_client=row.get("IsFreightOrderClient"),
        enterprise_vertical=row.get("EnterpriseVertical"),
        control_tower=row.get("ControlTower"),
        sales_source=row.get("SalesSource"),
        connectivity=row.get("Connectivity"),
        customer_segmentation=row.get("CustomerSegmentation"),
        raw_json=row,
    )


def resolve_country(db: Session, org_code: str, row: Dict[str, Any]) -> Tuple[str, str]:
    rules = db.query(CountryResolutionRule).filter(CountryResolutionRule.org_code == org_code).order_by(CountryResolutionRule.priority_order).all()
    if not rules:
        rules = [
            CountryResolutionRule(source_field="PickupFromCountryCode", lookup_type="DIRECT"),
            CountryResolutionRule(source_field="DeliveryToCountryCode", lookup_type="DIRECT"),
            CountryResolutionRule(source_field="DischargePort", lookup_type="UNLOCODE"),
            CountryResolutionRule(source_field="LastDischarge", lookup_type="UNLOCODE"),
            CountryResolutionRule(source_field="ConsigneeCountryCode", lookup_type="DIRECT"),
        ]
    for rule in rules:
        value = row.get(rule.source_field)
        if not value and rule.source_field == "ConsignorCountryCode":
            value = (row.get("ConsignorAddress") or {}).get("CountryCode")
        if not value and rule.source_field == "ConsigneeCountryCode":
            value = (row.get("ConsigneeAddress") or {}).get("CountryCode")
        if not value:
            continue
        lookup_type = rule.lookup_type.upper()
        if lookup_type == "DIRECT":
            return str(value).upper(), rule.source_field
        if lookup_type == "UNLOCODE":
            rec = db.query(UNLOCODE).filter(UNLOCODE.unlocode == value).first()
            if rec:
                return rec.country_code.upper(), rule.source_field
            if len(str(value)) >= 2:
                return str(value)[:2].upper(), rule.source_field
        if lookup_type == "LOCATION_NAME":
            rec = db.query(LocationCountryMapping).filter(LocationCountryMapping.location_name.ilike(str(value))).first()
            if rec:
                return rec.country_code.upper(), rule.source_field
    return "Unknown", "Unresolved"


def make_shipment(db: Session, process_number: int, org_code: str, row: Dict[str, Any]) -> ShipmentData:
    country, source = resolve_country(db, org_code, row)
    booking = row.get("Booking") or {}
    consignor = row.get("ConsignorAddress") or {}
    consignee = row.get("ConsigneeAddress") or {}
    return ShipmentData(
        process_number=process_number,
        org_code=org_code,
        shipment_id=row.get("ID"),
        booking_number=row.get("BookingNumber") or booking.get("BookingNumber"),
        booking_type=booking.get("Type"),
        booking_type_name=booking.get("TypeName"),
        booking_id=row.get("BookingID"),
        incoterm=row.get("Incoterm"),
        unique_consign_ref=row.get("UniqueConsignRef"),
        date_created=parse_dt(row.get("DateCreated")),
        transport_mode=row.get("TransportMode"),
        consignor_name=row.get("ConsignorName"),
        consignor_country=consignor.get("Country"),
        consignor_country_code=consignor.get("CountryCode"),
        pickup_name=row.get("PickupName"),
        pickup_from_country=row.get("PickupFromCountry"),
        pickup_from_country_code=row.get("PickupFromCountryCode"),
        pickup_from_city=row.get("PickupFromCity"),
        origin_port=row.get("OriginPort"),
        first_load=row.get("FirstLoad"),
        consignee_name=row.get("ConsigneeName"),
        consignee_country=consignee.get("Country"),
        consignee_country_code=consignee.get("CountryCode"),
        delivery_name=row.get("DeliveryName"),
        delivery_to_country=row.get("DeliveryToCountry"),
        delivery_to_country_code=row.get("DeliveryToCountryCode"),
        delivery_to_city=row.get("DeliveryToCity"),
        destination_port=row.get("DestinationPort"),
        discharge_port=row.get("DischargePort"),
        last_discharge=row.get("LastDischarge"),
        style_template=row.get("StyleTemplate"),
        resolved_country_code=country,
        resolved_country_source=source,
        raw_json=row,
    )


async def collect_invoice_data(db: Session, from_date: date, to_date: date) -> InvoiceProcess:
    process = InvoiceProcess(from_date=from_date, to_date=to_date, status="Running")
    db.add(process)
    db.commit()
    db.refresh(process)
    try:
        usage_rows = await read_usage(from_date, to_date, 100)
        for row in usage_rows:
            if not row.get("OrgCode"):
                continue
            upsert_organisation(db, row.get("OrgCode"), row.get("OrgFullName"))
            db.add(make_usage(process.process_number, row))
        db.commit()

        orgs = db.query(Organisation).filter(Organisation.get_shipment_data == True, Organisation.excluded == False).all()
        for org in orgs:
            shipment_rows = await read_shipments(org.org_code, from_date, to_date, 100)
            for row in shipment_rows:
                db.add(make_shipment(db, process.process_number, org.org_code, row))
            usage = db.query(ScmUsage).filter(ScmUsage.process_number == process.process_number, ScmUsage.org_code == org.org_code).first()
            if usage and usage.shipment_count != len(shipment_rows):
                db.add(InvoiceException(
                    process_number=process.process_number,
                    org_code=org.org_code,
                    exception_type="SHIPMENT_COUNT_DISCREPANCY",
                    severity="Warning",
                    message=f"OrganizationUsage_Read ShipmentCount is {usage.shipment_count}, but Shipments_Read returned {len(shipment_rows)} records. Final invoice will use Shipments_Read count."
                ))
        process.status = "Completed"
        process.completed_at = datetime.utcnow()
        db.commit()
        db.refresh(process)
        return process
    except Exception as exc:
        db.query(ScmUsage).filter(ScmUsage.process_number == process.process_number).delete()
        db.query(ShipmentData).filter(ShipmentData.process_number == process.process_number).delete()
        db.query(InvoiceException).filter(InvoiceException.process_number == process.process_number).delete()
        process.status = "Failed"
        process.error_message = str(exc)
        process.completed_at = datetime.utcnow()
        db.commit()
        raise


def calculate_tier_charge(quantity: int, tiers: List[ChargePlanRate]) -> Tuple[Decimal, Decimal]:
    total = Decimal("0")
    first_rate = Decimal("0")
    for tier in sorted(tiers, key=lambda t: t.from_quantity):
        if quantity < tier.from_quantity:
            continue
        upper = tier.to_quantity if tier.to_quantity is not None else quantity
        chargeable_qty = min(quantity, upper) - tier.from_quantity + 1
        if chargeable_qty > 0:
            rate = Decimal(tier.unit_price)
            if first_rate == 0:
                first_rate = rate
            total += Decimal(chargeable_qty) * rate
    return total, first_rate


def create_final_invoice(db: Session, process_number: int) -> List[FinalInvoiceLine]:
    db.query(FinalInvoiceLine).filter(FinalInvoiceLine.process_number == process_number).delete()
    usages = db.query(ScmUsage).filter(ScmUsage.process_number == process_number).all()
    lines: List[FinalInvoiceLine] = []
    for usage in usages:
        org = db.query(Organisation).filter(Organisation.org_code == usage.org_code).first()
        if org and org.excluded:
            continue
        charge_plan_code = (org.charge_plan_code if org and org.charge_plan_code else "DEFAULT_2026")
        plan = db.query(ChargePlan).filter(ChargePlan.charge_plan_code == charge_plan_code).first()
        tiers = db.query(ChargePlanRate).filter(
            ChargePlanRate.charge_plan_code == charge_plan_code,
            ChargePlanRate.metric_code == "SHIPMENT",
            ChargePlanRate.is_active == True,
        ).all()
        currency = plan.currency_code if plan else "EUR"
        if org and org.get_shipment_data:
            grouped = defaultdict(int)
            shipments = db.query(ShipmentData).filter(ShipmentData.process_number == process_number, ShipmentData.org_code == usage.org_code).all()
            for sh in shipments:
                grouped[sh.resolved_country_code or "Unknown"] += 1
            for country, qty in grouped.items():
                total, unit = calculate_tier_charge(qty, tiers)
                line = FinalInvoiceLine(process_number=process_number, org_code=usage.org_code, org_full_name=usage.org_full_name, country_code=country, division=org.division if org else None, metric_code="SHIPMENT", quantity=qty, unit_price=unit, total_cost=total, currency_code=currency, source="Shipments_Read")
                db.add(line); lines.append(line)
        else:
            qty = usage.shipment_count or 0
            total, unit = calculate_tier_charge(qty, tiers)
            line = FinalInvoiceLine(process_number=process_number, org_code=usage.org_code, org_full_name=usage.org_full_name, country_code=org.country_code if org else None, division=org.division if org else None, metric_code="SHIPMENT", quantity=qty, unit_price=unit, total_cost=total, currency_code=currency, source="OrganizationUsage_Read")
            db.add(line); lines.append(line)
    db.commit()
    return lines


def _write_sheet(ws, rows: List[dict]):
    if not rows:
        ws.append(["No data"]); return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    for idx, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = 18


def generate_excel(db: Session, process_number: int) -> str:
    GENERATED_DIR.mkdir(exist_ok=True)
    create_final_invoice(db, process_number)
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Summary"
    final_rows = db.query(FinalInvoiceLine).filter(FinalInvoiceLine.process_number == process_number).all()
    _write_sheet(ws, [{"OrgCode": r.org_code, "OrgFullName": r.org_full_name, "CountryCode": r.country_code, "Division": r.division, "Metric": r.metric_code, "Quantity": r.quantity, "UnitPrice": float(r.unit_price), "TotalCost": float(r.total_cost), "Currency": r.currency_code, "Source": r.source} for r in final_rows])
    ws = wb.create_sheet("SCM Usage")
    usages = db.query(ScmUsage).filter(ScmUsage.process_number == process_number).all()
    _write_sheet(ws, [{"OrgCode": u.org_code, "OrgFullName": u.org_full_name, "ShipmentCount": u.shipment_count, "OrderCount": u.order_count, "BookingCount": u.booking_count, "YearMonth": u.year_month} for u in usages])
    ws = wb.create_sheet("Shipment Data")
    shipments = db.query(ShipmentData).filter(ShipmentData.process_number == process_number).all()
    _write_sheet(ws, [{"OrgCode": s.org_code, "UniqueConsignRef": s.unique_consign_ref, "DateCreated": s.date_created, "TransportMode": s.transport_mode, "ResolvedCountryCode": s.resolved_country_code, "ResolvedCountrySource": s.resolved_country_source, "PickupFromCountryCode": s.pickup_from_country_code, "DeliveryToCountryCode": s.delivery_to_country_code, "DischargePort": s.discharge_port, "LastDischarge": s.last_discharge} for s in shipments])
    ws = wb.create_sheet("Exceptions")
    exceptions = db.query(InvoiceException).filter(InvoiceException.process_number == process_number).all()
    _write_sheet(ws, [{"OrgCode": e.org_code, "Type": e.exception_type, "Severity": e.severity, "Message": e.message, "Resolved": e.is_resolved} for e in exceptions])
    path = GENERATED_DIR / f"edc_invoice_process_{process_number}.xlsx"
    wb.save(path)
    process = db.query(InvoiceProcess).filter(InvoiceProcess.process_number == process_number).first()
    if process:
        process.excel_file_path = str(path)
        db.commit()
    return str(path)
